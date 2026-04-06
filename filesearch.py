import tkinter as tk
from tkinter import filedialog, ttk, messagebox
import os
import csv
import threading
import queue

# Import required libraries for content searching
try:
    import PyPDF2
    import docx
    import openpyxl
except ImportError:
    messagebox.showerror("Missing Libraries", "Please install required libraries: pip install PyPDF2 python-docx openpyxl")
    exit()


class FileSearchApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Advanced File Search")
        self.found_files = [] 
        self.search_thread = None
        self.cancel_search_flag = False

        # --- Directory Selection ---
        self.dir_frame = ttk.LabelFrame(root, text="1. Select a folder to search")
        self.dir_frame.pack(padx=10, pady=10, fill="x")
        self.browse_button = ttk.Button(self.dir_frame, text="Browse...", command=self.browse_directory)
        self.browse_button.pack(side="left", padx=5, pady=5)
        self.selected_dir_label = ttk.Label(self.dir_frame, text="No folder selected")
        self.selected_dir_label.pack(side="left", padx=5, pady=5)

        # --- Search Criteria ---
        self.search_frame = ttk.LabelFrame(root, text="2. Define search criteria")
        self.search_frame.pack(padx=10, pady=5, fill="x")

        ttk.Label(self.search_frame, text="Search for:").pack(side="left", padx=5, pady=5)
        self.search_entry = ttk.Entry(self.search_frame, width=30)
        self.search_entry.pack(side="left", padx=5, pady=5, fill="x", expand=True)

        self.search_type_var = tk.StringVar(value="filename")
        ttk.Radiobutton(self.search_frame, text="File Name", variable=self.search_type_var, value="filename").pack(side="left", padx=5)
        ttk.Radiobutton(self.search_frame, text="File Content", variable=self.search_type_var, value="content").pack(side="left", padx=5)

        # --- File Type Selection ---
        self.types_frame = ttk.LabelFrame(root, text="3. Select file types")
        self.types_frame.pack(padx=10, pady=5, fill="x")
        self.file_type_vars = {}
        
        self.file_types = {
            "All Files": "*",
            "Office Documents": ['.docx', '.xlsx', '.pptx', '.doc', '.xls', '.ppt'],
            "PDF Files": ['.pdf'],
            "Text Files": ['.txt', '.rtf', '.csv', '.json', '.xml', '.md'],
            "Image Files": ['.jpg', '.jpeg', '.png', '.gif', '.bmp'],
            "Video/Audio Files": ['.mp4', '.mov', '.avi', '.mp3', '.wav'],
            "Scripts & Code": ['.py', '.js', '.html', '.css', '.java'],
        }

        for i, (label, exts) in enumerate(self.file_types.items()):
            var = tk.BooleanVar()
            cb = ttk.Checkbutton(self.types_frame, text=label, variable=var, command=self.on_file_type_select)
            cb.grid(row=i // 3, column=i % 3, sticky="w", padx=5, pady=2)
            self.file_type_vars[label] = var
        
        self.file_type_vars["All Files"].set(True)
        self.on_file_type_select()

        # --- Action Buttons ---
        self.action_frame = ttk.Frame(root)
        self.action_frame.pack(padx=10, pady=10, fill="x")

        self.search_button = ttk.Button(self.action_frame, text="Search", command=self.start_search)
        self.search_button.pack(side="left", padx=(0, 5))
        
        self.cancel_button = ttk.Button(self.action_frame, text="Cancel", command=self.cancel_search, state="disabled")
        self.cancel_button.pack(side="left", padx=(0, 10))

        self.export_button = ttk.Button(self.action_frame, text="Export Results...", command=self.export_results, state="disabled")
        self.export_button.pack(side="right", padx=5)
        
        # --- Results Display ---
        self.results_frame = ttk.LabelFrame(root, text="Results")
        self.results_frame.pack(padx=10, pady=10, fill="both", expand=True)

        self.results_listbox = tk.Listbox(self.results_frame)
        self.results_listbox.pack(side="left", fill="both", expand=True, padx=5, pady=5)
        self.scrollbar = ttk.Scrollbar(self.results_frame, orient="vertical", command=self.results_listbox.yview)
        self.scrollbar.pack(side="right", fill="y")
        self.results_listbox.config(yscrollcommand=self.scrollbar.set)
        
    def browse_directory(self):
        directory = filedialog.askdirectory()
        if directory:
            self.directory = directory
            self.selected_dir_label.config(text=directory)

    def on_file_type_select(self):
        is_all_checked = self.file_type_vars["All Files"].get()
        state = "disabled" if is_all_checked else "normal"
        for label, var in self.file_type_vars.items():
            if label != "All Files":
                for child in self.types_frame.winfo_children():
                    if child.cget("text") == label:
                        child.config(state=state)

    def get_selected_extensions(self):
        if self.file_type_vars["All Files"].get():
            return ["*"]
        selected_extensions = [ext for label, var in self.file_type_vars.items() if var.get() for ext in self.file_types.get(label, [])]
        return list(set(selected_extensions))

    def start_search(self):
        if not hasattr(self, 'directory'):
            messagebox.showwarning("No Directory", "Please select a directory first.")
            return
        keyword = self.search_entry.get()
        if not keyword:
            messagebox.showwarning("No Keyword", "Please enter a search term.")
            return

        self.search_button.config(state="disabled")
        self.cancel_button.config(state="normal")
        self.export_button.config(state="disabled")
        self.results_listbox.delete(0, tk.END)
        self.results_listbox.insert(tk.END, "Searching, please wait...")
        self.found_files.clear()
        self.cancel_search_flag = False

        self.result_queue = queue.Queue()
        self.search_thread = threading.Thread(
            target=self.search_worker,
            args=(keyword.lower(), self.get_selected_extensions(), self.search_type_var.get())
        )
        self.search_thread.start()
        self.root.after(100, self.process_queue)

    def cancel_search(self):
        if self.search_thread and self.search_thread.is_alive():
            self.cancel_search_flag = True
            self.cancel_button.config(text="Cancelling...", state="disabled")

    def process_queue(self):
        try:
            while True:
                msg = self.result_queue.get_nowait()
                if msg is None: # Sentinel value indicates thread is done
                    self.search_finished()
                    return
                elif msg == "CANCELLED":
                    self.search_finished(cancelled=True)
                    return
                else:
                    if self.results_listbox.get(0) == "Searching, please wait...":
                        self.results_listbox.delete(0, tk.END)
                    self.results_listbox.insert(tk.END, msg)
                    self.found_files.append(msg)
        except queue.Empty:
            self.root.after(100, self.process_queue) # Check again later

    def search_finished(self, cancelled=False):
        self.search_button.config(state="normal")
        self.cancel_button.config(text="Cancel", state="disabled")
        if self.found_files:
            self.export_button.config(state="normal")
        
        if cancelled:
            if not self.found_files: # If no results were found before cancelling
                self.results_listbox.delete(0, tk.END)
            self.results_listbox.insert(tk.END, "--- Search Cancelled ---")
        elif not self.found_files: # Search finished with no results
             self.results_listbox.delete(0, tk.END)
             self.results_listbox.insert(tk.END, "No matching files found.")


    def search_worker(self, keyword, extensions, search_mode):
        for root_dir, _, files in os.walk(self.directory, topdown=True):
            if self.cancel_search_flag:
                self.result_queue.put("CANCELLED")
                return
            for file in files:
                if self.cancel_search_flag:
                    self.result_queue.put("CANCELLED")
                    return
                
                file_lower = file.lower()
                file_path = os.path.join(root_dir, file)

                if extensions[0] != "*" and not any(file_lower.endswith(ext) for ext in extensions):
                    continue

                if search_mode == 'filename':
                    if keyword in file_lower:
                        self.result_queue.put(file_path)
                elif search_mode == 'content':
                    try:
                        content = self._read_file_content(file_path)
                        if content is not None and keyword in content.lower():
                            self.result_queue.put(file_path)
                    except Exception as e:
                        print(f"Could not read {file_path}: {e}")
        
        self.result_queue.put(None) # Signal completion

    def _read_file_content(self, file_path):
        _, ext = os.path.splitext(file_path.lower())
        content = ""
        try:
            if ext == '.pdf':
                with open(file_path, 'rb') as f:
                    reader = PyPDF2.PdfReader(f, strict=False)
                    for page in reader.pages:
                        content += page.extract_text() or ""
            elif ext == '.docx':
                doc = docx.Document(file_path)
                content = "\n".join([para.text for para in doc.paragraphs])
            elif ext == '.xlsx':
                workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                for sheet in workbook:
                    for row in sheet.iter_rows():
                        for cell in row:
                            if cell.value:
                                content += str(cell.value) + " "
            else: # Treat as plain text
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    content = f.read()
        except Exception:
            return None # Ignore files that can't be read
        return content

    def export_results(self):
        if not self.found_files:
            messagebox.showinfo("No Results", "There are no results to export.")
            return
        file_path = filedialog.asksaveasfilename(
            defaultextension=".csv",
            filetypes=[("CSV file", "*.csv"), ("Text file", "*.txt")]
        )
        if not file_path: return
        try:
            if file_path.endswith('.csv'):
                with open(file_path, 'w', newline='', encoding='utf-8') as f:
                    writer = csv.writer(f)
                    writer.writerow(["File Path"])
                    writer.writerows([[path] for path in self.found_files])
            else:
                with open(file_path, 'w', encoding='utf-8') as f:
                    f.write('\n'.join(self.found_files))
            messagebox.showinfo("Export Successful", f"Results successfully exported to {os.path.basename(file_path)}")
        except Exception as e:
            messagebox.showerror("Export Error", f"An error occurred during export: {e}")

if __name__ == "__main__":
    root = tk.Tk()
    app = FileSearchApp(root)
    root.mainloop()
