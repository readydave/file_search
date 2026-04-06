# app/search/search_worker.py

"""
Contains the core search logic that runs in a background thread
to prevent the UI from freezing.
"""

import os
import threading

# Import required libraries for content searching
try:
    import PyPDF2
    import docx
    import openpyxl
except ImportError:
    # This is a fallback; the main UI should handle this more gracefully.
    print("Error: Missing required libraries. Please run 'pip install PyPDF2 python-docx openpyxl'")


class SearchWorker:
    """
    This class performs the file search in a separate thread.
    """
    def __init__(self, directory, keyword, extensions, search_mode, result_queue, cancel_event):
        self.directory = directory
        self.keyword = keyword.lower()
        self.extensions = extensions
        self.search_mode = search_mode
        self.result_queue = result_queue
        self.cancel_event = cancel_event

    def run(self):
        """The main worker method that walks the directory tree and searches files."""
        try:
            for root_dir, _, files in os.walk(self.directory, topdown=True):
                if self.cancel_event.is_set():
                    break
                for file in files:
                    if self.cancel_event.is_set():
                        break
                    
                    file_lower = file.lower()
                    file_path = os.path.join(root_dir, file)

                    # Filter by file extension if 'All Files' is not selected
                    if self.extensions[0] != "*" and not any(file_lower.endswith(ext) for ext in self.extensions):
                        continue

                    # Perform the search
                    if self.search_mode == 'filename':
                        if self.keyword in file_lower:
                            self.result_queue.put(file_path)
                    elif self.search_mode == 'content':
                        try:
                            content = self._read_file_content(file_path)
                            if content is not None and self.keyword in content.lower():
                                self.result_queue.put(file_path)
                        except Exception as e:
                            # Log errors to console but don't stop the search
                            print(f"Could not read {file_path}: {e}")
        finally:
            # Use a sentinel value to signal that the search is complete.
            # Check if cancellation was the reason for stopping.
            if self.cancel_event.is_set():
                self.result_queue.put("CANCELLED")
            else:
                self.result_queue.put("FINISHED")

    def _read_file_content(self, file_path):
        """Helper function to read content from various file types."""
        _, ext = os.path.splitext(file_path.lower())
        content = ""
        try:
            if ext == '.pdf':
                with open(file_path, 'rb') as f:
                    # PyPDF2 can be strict; some PDFs might cause issues.
                    reader = PyPDF2.PdfReader(f, strict=False)
                    for page in reader.pages:
                        content += page.extract_text() or ""
            elif ext == '.docx':
                doc = docx.Document(file_path)
                content = "\n".join([para.text for para in doc.paragraphs])
            elif ext == '.xlsx':
                # data_only=True reads cell values, not formulas
                workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                for sheet in workbook:
                    for row in sheet.iter_rows():
                        for cell in row:
                            if cell.value:
                                content += str(cell.value) + " "
            else:  # Treat as plain text by default
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    content = f.read()
        except Exception:
            # This can happen with binary files or protected files.
            # Return None to indicate we couldn't read it.
            return None
        return content

